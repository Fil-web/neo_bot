import csv
from io import StringIO
from pathlib import Path
from typing import Optional

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


class DocumentParsingError(Exception):
    pass


class DocumentParser:
    def __init__(self, max_chars: int) -> None:
        self._max_chars = max_chars

    def parse(
        self,
        file_path: Path,
        file_name: Optional[str] = None,
        mime_type: Optional[str] = None,
    ) -> str:
        suffix = (Path(file_name).suffix if file_name else file_path.suffix).lower()

        if suffix in {".txt", ".md", ".log", ".json", ".py", ".js", ".ts", ".html", ".css"}:
            text = file_path.read_text(encoding="utf-8", errors="ignore")
            return self._truncate(text)
        if suffix == ".csv":
            return self._truncate(self._parse_csv(file_path))
        if suffix == ".pdf":
            return self._truncate(self._parse_pdf(file_path))
        if suffix == ".docx":
            return self._truncate(self._parse_docx(file_path))
        if suffix in {".xlsx", ".xlsm"}:
            return self._truncate(self._parse_xlsx(file_path))

        if mime_type and mime_type.startswith("text/"):
            text = file_path.read_text(encoding="utf-8", errors="ignore")
            return self._truncate(text)

        raise DocumentParsingError(f"Unsupported document type: {suffix or mime_type or 'unknown'}")

    def _truncate(self, text: str) -> str:
        cleaned = text.strip()
        if not cleaned:
            raise DocumentParsingError("Document is empty")
        if len(cleaned) <= self._max_chars:
            return cleaned
        return cleaned[: self._max_chars] + "\n\n[Фрагмент обрезан по лимиту]"

    def _parse_pdf(self, file_path: Path) -> str:
        reader = PdfReader(str(file_path))
        parts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                parts.append(f"[Страница {index}]\n{page_text}")
        return "\n\n".join(parts)

    def _parse_docx(self, file_path: Path) -> str:
        document = DocxDocument(str(file_path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n".join(paragraphs)

    def _parse_csv(self, file_path: Path) -> str:
        content = file_path.read_text(encoding="utf-8", errors="ignore")
        reader = csv.reader(StringIO(content))
        rows: list[str] = []
        for index, row in enumerate(reader, start=1):
            rows.append(f"Строка {index}: {' | '.join(cell.strip() for cell in row)}")
            if index >= 100:
                break
        return "\n".join(rows)

    def _parse_xlsx(self, file_path: Path) -> str:
        workbook = load_workbook(filename=str(file_path), read_only=True, data_only=True)
        chunks: list[str] = []
        for sheet in workbook.worksheets[:5]:
            chunks.append(f"[Лист: {sheet.title}]")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    chunks.append(f"Строка {row_index}: {' | '.join(values)}")
                if row_index >= 50:
                    break
        return "\n".join(chunks)
